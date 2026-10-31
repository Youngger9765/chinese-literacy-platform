#!/usr/bin/env python3
"""教授節點 → 108 課綱三層 → 179 課繼承。重建 lesson-indicators.json 的節點欄位與審核表。

用法（從 repo 根目錄或任何地方，路徑都是絕對）：
  python3 build_lesson_indicators.py                      # 線上抓 API，寫回 json＋審核表
  python3 build_lesson_indicators.py --stories-json X     # 用存好的 API 回應（可重現）
  python3 build_lesson_indicators.py --check              # 只算不寫，印統計

冪等：會先拿掉自己上次加的欄位再重算，所以跑兩次結果相同（xlsx 只比儲存格內容）。
原本的 AI 逐課欄位（primary_indicator / skills / core_competencies / learning_contents）不動，
只加 superseded_by_node 標記。
"""
from __future__ import annotations

import argparse
import collections
import difflib
import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
REPO = SKILL_DIR.parent.parent.parent
DEFAULTS = {
    "xlsx": Path("/Users/young/project/kist-curriculum/教材原檔/自學教材總表0816.xlsx"),
    "nodes": Path("/Users/young/project/kist-curriculum/data/node_definitions.json"),
    "curriculum": Path.home() / ".claude/skills/tw-108-curriculum/data/chinese.json",
    "zonggang": Path.home() / ".claude/skills/tw-108-curriculum/data/zonggang.json",
    "suggestions": SKILL_DIR / "node_code_suggestions.json",
    "lesson_json": REPO / "docs/design/reading-babies/lesson-indicators.json",
    "review_xlsx": REPO / "docs/design/reading-babies/node-indicator-review.xlsx",
    "code_csv": REPO / "docs/curriculum/2026-08-14-second-edition-lesson-code-correction.csv",
    "api": "https://lingoleap-backend-issue-3780-oja2sffiya-de.a.run.app/api/stories?page_size=300",
}
STAGE_OF_GRADE = {4: "Ⅱ", 5: "Ⅲ", 6: "Ⅲ", 7: "Ⅳ", 8: "Ⅳ", 9: "Ⅳ"}
SCHOOL_OF_STAGE = {"Ⅱ": "E", "Ⅲ": "E", "Ⅳ": "J"}
ADDED_LESSON_KEYS = ["professor_row", "professor_strategy", "professor_node", "mapping_source",
                     "node_indicators", "superseded_by_node"]
ADDED_TOP_KEYS = ["professor_source", "node_mapping", "node_codes", "issues_19", "node_mapping_coverage"]
SRC_RANK = {"teacher_reviewed": 0, "professor_table": 1, "rule_from_professor_field": 2,
            "professor_ai_map": 3, "ai_inferred": 4}

# 教授 F 欄策略「──」前的類別名 → 家族（驗證節點歸屬、挑主節點用）
FAMILY_KEYWORDS = [
    ("F01", "摘要|上位概念|歸納概念"), ("F02", "推論"), ("F03", "圖表|表格|文圖"), ("F04", "提問"),
    ("F05", "解題|監控|調整閱讀|轉折|標點|帶著問題"), ("F06", "詞彙|詞義|拆詞|語詞"), ("F07", "句型|複句"),
    ("F08", "寫作手法"), ("F09", "科學推理"), ("F10", "媒體"), ("F11", "多文本"), ("F12", "反思|4F"),
    ("F13", "文言文"), ("F14", "品格|生涯|SEL"),
]
# R1 文體 → 學習內容文本表述類別
GENRE_TO_CATEGORY = {
    "記敘文": ["Ba"], "抒情文": ["Bb"], "說明文": ["Bc"], "説明文": ["Bc"], "議論文": ["Bd"],
    "應用文": ["Be"], "記敘抒情文": ["Ba", "Bb"], "記敘說理文": ["Ba", "Bd"],
    "說明/議論文": ["Bc", "Bd"], "科普記敘文": ["Ba", "Bc"],
}
STRUCTURE_CODE = {  # 該階段有「X 文本的結構」條目就用它，否則用該類第 1 條（並標註）
    ("Ba", "Ⅱ"): "Ba-Ⅱ-1", ("Bb", "Ⅱ"): "Bb-Ⅱ-6", ("Bc", "Ⅲ"): "Bc-Ⅲ-4",
    ("Bd", "Ⅲ"): "Bd-Ⅲ-3", ("Be", "Ⅱ"): "Be-Ⅱ-4",
}
VOCAB_CODES = {"Ⅱ": ["Ab-Ⅱ-5", "Ab-Ⅱ-6"], "Ⅲ": ["Ab-Ⅲ-5", "Ab-Ⅲ-6"], "Ⅳ": ["Ab-Ⅳ-4", "Ab-Ⅳ-5"]}
CLASSICAL_CODES = ["Ab-Ⅳ-6", "Ab-Ⅳ-7"]
# R4 教授「議題」欄關鍵字 → 19 項議題（只收字面或明確同義，其餘不對應、列出來）
ISSUE_KEYWORDS = [
    ("品德", r"品格|品德|堅毅|堅持|自律|行善|仁慈|同理心|感恩|樂觀|助人|運動家精神|毅力|克服挫折|正向思考|自我控制"),
    ("環境", r"環境|永續|氣候|自然之美"),
    ("性別平等", r"性別"),
    ("多元文化", r"多元文化"),
    ("科技", r"科技"),
    ("閱讀素養", r"閱讀|讀書"),
    ("家庭教育", r"代間"),
]


def clean(s) -> str | None:
    if s is None:
        return None
    s = "".join(ch for ch in str(s) if not (0xE0100 <= ord(ch) <= 0xE01EF or 0xFE00 <= ord(ch) <= 0xFE0F))
    s = s.strip()
    return s or None


def norm(s) -> str:
    s = unicodedata.normalize("NFKC", clean(s) or "")
    s = s.replace("臺", "台").replace("没", "沒").replace("燬", "毀")
    return re.sub(r"[\s：:？?！!，,、「」『』\"“”\-—–─.。（）()＃#…]", "", s)


# ---------- 讀教授總表 ----------
def load_professor(xlsx: Path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows = []

    def words(r, start):
        return [w for w in (clean(x) for x in r[start:]) if w and not w.startswith("http")]

    for i, r in enumerate(wb["1.總表"].iter_rows(min_row=2, values_only=True), start=2):
        if not clean(r[2]):
            continue
        g, n = int(r[0]), int(r[1])
        rows.append(dict(sheet="1.總表", excel_row=i, slot=f"G{g}-L{n}", grade=g, title=clean(r[2]),
                         genre=clean(r[3]), type=clean(r[4]), strategy=clean(r[5]), issue=clean(r[6]),
                         sixth_title=clean(r[7]), text_type=clean(r[8]), words=words(r, 9)))
    for i, r in enumerate(wb["2.體育生的品格教材"].iter_rows(min_row=2, values_only=True), start=2):
        m = re.match(r"體-(\d+)$", clean(r[0]) or "")
        if not m:
            continue
        rows.append(dict(sheet="2.體育生的品格教材", excel_row=i, slot=f"體-{int(m.group(1))}", grade=int(r[2]),
                         title=clean(r[3]), genre=clean(r[4]), type=clean(r[5]), strategy=clean(r[6]),
                         issue=None, sixth_title=clean(r[7]), text_type=clean(r[8]), words=words(r, 9)))
    for i, r in enumerate(wb["3.文言文"].iter_rows(min_row=2, values_only=True), start=2):
        m = re.match(r"文-(\d+)$", clean(r[0]) or "")
        if not m:
            continue
        rows.append(dict(sheet="3.文言文", excel_row=i, slot=f"文-{int(m.group(1))}", grade=int(r[1]),
                         title=clean(r[2]), genre=clean(r[3]), type=clean(r[4]), strategy=clean(r[6]),
                         issue=clean(r[5]), sixth_title=clean(r[7]), text_type=clean(r[8]), words=words(r, 9)))
    by_slot = {}
    for row in rows:
        assert row["slot"] not in by_slot, f"總表課號重複 {row['slot']}"
        by_slot[row["slot"]] = row
    return by_slot


def api_slot(grade_code: str) -> str | None:
    m = re.match(r"(G\d)-L(\d+)$", grade_code)
    if m:
        return f"{m.group(1)}-L{int(m.group(2))}"
    m = re.match(r"(體|文)-L(\d+)$", grade_code)
    return f"{m.group(1)}-{int(m.group(2))}" if m else None


# ---------- 課綱原文 ----------
def load_codes(curriculum: Path, existing: dict):
    data = json.load(open(curriculum, encoding="utf-8"))
    k12_url = existing["source"]["url"]
    k12 = {}
    for sec in ("indicators", "learning_contents", "core_competencies"):
        for code, v in existing.get(sec, {}).items():
            k12[code.replace("-I-", "-Ⅰ-")] = v
    codes, offsets = {}, set()
    for e in data["entries"]:
        rec = dict(code=e["code"], type=e["type"], text=e["text"], naer_pdf=e["source_pdf"], naer_page=e["page"])
        old = k12.get(e["code"])
        if old:
            offsets.add(old["pdf_page"] - e["page"])
            rec["k12ea_copy"] = "一致" if old["text"].strip() == e["text"].strip() else "k12ea 抄本有誤，已用國教院原文"
        else:
            rec["k12ea_copy"] = "k12ea 抄本未收"
        codes[e["code"]] = rec
    assert offsets == {2}, f"k12ea 與國教院 PDF 頁碼差不是固定 2：{offsets}"
    for rec in codes.values():
        rec["k12ea_page"] = rec["naer_page"] + 2
        rec["k12ea_url"] = f"{k12_url}#page={rec['k12ea_page']}"
    return codes


def code_ref(codes, code, **extra):
    c = codes.get(code)
    if not c:
        raise SystemExit(f"官方課綱沒有這個代碼：{code}")
    return dict(code=code, text=c["text"], pdf_page=c["k12ea_page"], url=c["k12ea_url"], **extra)


# ---------- 節點建議 ----------
def resolve_node(sug, node, stages):
    fam = sug["family_defaults"][node["family"]]
    own = sug["nodes"].get(node["id"], {})
    out = {"perf": {}, "content": {}, "core": own.get("core", fam["core"]),
           "why": "；".join(x for x in (fam.get("why"), own.get("why")) if x),
           "source": "teacher_reviewed" if own.get("reviewed_by") else "ai_inferred"}
    for layer in ("perf", "content"):
        for st in stages:
            val = (own.get(layer) or {}).get(st) or (fam.get(layer) or {}).get(st)
            if not val:
                raise SystemExit(f"{node['id']} 缺 {layer} 第{st}學習階段")
            lo, hi = (1, 2) if layer == "perf" else (1, 3)
            assert lo <= len(val) <= hi, f"{node['id']} {layer} {st} 條數 {len(val)}"
            out[layer][st] = val
    return out


def families_of_strategy(strategy: str | None) -> set:
    """F 欄「類別──細項」只看類別段；類別段認不出再看全文。回傳可能的家族集合（空＝認不出）。"""
    if not strategy:
        return set()
    head = re.split(r"──|─|－", strategy)[0]
    for text in (head, strategy):
        hits = {fid for fid, pat in FAMILY_KEYWORDS if re.search(pat, text)}
        if hits:
            return hits
    return set()


def issues_from(raw: str | None, field: str):
    out = []
    if not raw or raw in ("無", "None"):
        return out
    for issue, pat in ISSUE_KEYWORDS:
        if re.search(pat, raw):
            out.append(dict(issue=issue, raw=raw, from_field=field, source="rule_from_professor_field",
                            rule=f"R4 {field}關鍵字 /{pat}/ → {issue}"))
    return out


def genre_contents(codes, genre, typ, stage):
    out, notes = [], []
    if genre == "文言文":
        out.append(code_ref(codes, "Ad-Ⅳ-4", source="rule_from_professor_field", rule="R3 文體＝文言文 → 非韻文古文"))
        return out, notes
    cats = GENRE_TO_CATEGORY.get(genre or "")
    if not cats:
        notes.append(f"文體「{genre}」沒有規則對應的文本表述類別")
        return out, notes
    for cat in cats:
        if cat == "Be" and typ == "讀書報告" and stage in ("Ⅱ", "Ⅲ", "Ⅳ"):
            code = f"Be-{stage}-3"
            out.append(code_ref(codes, code, source="rule_from_professor_field",
                                rule="R1 文體＝應用文且類型＝讀書報告 → 學習應用"))
            continue
        code = STRUCTURE_CODE.get((cat, stage))
        rule = f"R1 文體＝{genre} → {cat} 類，取『文本結構』條"
        if not code:
            first = sorted((c for c in codes if c.startswith(f"{cat}-{stage}-")),
                           key=lambda c: int(c.rsplit("-", 1)[1]))
            if not first:
                notes.append(f"第{stage}學習階段沒有 {cat} 類學習內容（文體「{genre}」無法對應）")
                continue
            code = first[0]
            rule = f"R1 文體＝{genre} → {cat} 類；該階段無『文本結構』條，取該類第 1 條"
        out.append(code_ref(codes, code, source="rule_from_professor_field", rule=rule))
    return out, notes


def merge_items(items):
    """同層同代碼只留一筆，來源取最強，其餘記在 also_from。"""
    by = collections.OrderedDict()
    for it in items:
        key = it.get("code") or it.get("issue")
        if key in by:
            keep = by[key]
            if SRC_RANK[it["source"]] < SRC_RANK[keep["source"]]:
                it.setdefault("also_from", []).append(keep.get("via") or keep.get("rule"))
                by[key] = it
            else:
                keep.setdefault("also_from", []).append(it.get("via") or it.get("rule"))
        else:
            by[key] = it
    return list(by.values())


def main():
    ap = argparse.ArgumentParser()
    for k in ("xlsx", "nodes", "curriculum", "zonggang", "suggestions", "lesson_json", "review_xlsx", "code_csv"):
        ap.add_argument(f"--{k.replace('_', '-')}", type=Path, default=DEFAULTS[k])
    ap.add_argument("--api", default=DEFAULTS["api"])
    ap.add_argument("--stories-json", type=Path, help="用存好的 /api/stories 回應，不連線")
    ap.add_argument("--save-stories", type=Path, help="把抓到的 API 回應存檔")
    ap.add_argument("--check", action="store_true", help="只算不寫")
    a = ap.parse_args()

    existing = json.load(open(a.lesson_json, encoding="utf-8"))
    for k in ADDED_TOP_KEYS:
        existing.pop(k, None)
    for les in existing["lessons"]:
        for k in ADDED_LESSON_KEYS:
            les.pop(k, None)

    if a.stories_json:
        api = json.load(open(a.stories_json, encoding="utf-8"))
    else:
        with urllib.request.urlopen(a.api, timeout=60) as r:
            api = json.loads(r.read())
        if a.save_stories:
            a.save_stories.write_text(json.dumps(api, ensure_ascii=False), encoding="utf-8")
    stories = api["stories"]
    assert len(stories) == api["total"], f"API 只拿到 {len(stories)}／{api['total']}"
    sid_set = {s["id"] for s in stories}
    assert sid_set == {l["story_id"] for l in existing["lessons"]}, "json 的課跟 API 不一致"

    prof = load_professor(a.xlsx)
    codes = load_codes(a.curriculum, existing)
    nd = json.load(open(a.nodes, encoding="utf-8"))
    fam_name = {f["id"]: f["name"] for f in nd["families"]}
    nodes = {n["id"]: n for n in nd["nodes"]}
    sug = json.load(open(a.suggestions, encoding="utf-8"))
    zg = json.load(open(a.zonggang, encoding="utf-8"))

    # ---- 1. 179 課對總表 ----
    match = {}
    unmatched = []
    for s in stories:
        slot = api_slot(s["grade_code"])
        row = prof.get(slot)
        if not row:
            unmatched.append(dict(story_id=s["id"], grade_code=s["grade_code"], title=s["title"], reason="總表沒有這個課號"))
            continue
        a_t, p_t = norm(s["title"]), norm(row["title"])
        ratio = difflib.SequenceMatcher(None, a_t, p_t).ratio()
        if a_t == p_t:
            kind = "課號相同、課名相同"
        elif a_t in p_t or p_t in a_t or ratio >= 0.6:
            kind = f"課號相同、課名有差異（相似度 {ratio:.2f}）"
        else:
            unmatched.append(dict(story_id=s["id"], grade_code=s["grade_code"], title=s["title"],
                                  reason=f"課號相同但課名不像（總表：{row['title']}，相似度 {ratio:.2f}）"))
            continue
        if s["grade"].isdigit() and int(s["grade"]) != row["grade"]:
            unmatched.append(dict(story_id=s["id"], grade_code=s["grade_code"], title=s["title"],
                                  reason=f"年級不一致（API {s['grade']}，總表 {row['grade']}）"))
            continue
        match[s["id"]] = dict(row=row, kind=kind, api_title=s["title"])

    # ---- 2. 每個節點用到哪些學習階段（含線上沒有的課）----
    # 2a. 教授地圖的每個課號逐條驗證：先當總表0816課號；F 欄家族不一致再試 08-14 校正清單的舊課號
    old_title = {}
    import csv
    for r in csv.DictReader(open(a.code_csv, encoding="utf-8")):
        if r["sheet"] != "(線上)":
            m_ = re.match(r"(G\d)-L0*(\d+)$", r["new_code"])
            k = f"{m_.group(1)}-L{m_.group(2)}" if m_ else re.sub(r"^(體|文)-L?0*", r"\1-", r["new_code"])
            old_title[k] = r["title"]
    title_slots = collections.defaultdict(list)
    for k, r in prof.items():
        if not k.endswith("-L0"):
            title_slots[norm(r["title"])].append(k)

    def fits(slot, fam):
        fs = families_of_strategy(prof[slot]["strategy"]) if slot in prof else set()
        return (not fs) or fam in fs

    def translate(c, fam):
        t = old_title.get(c)
        cands = [x for x in title_slots.get(norm(t), []) if x.split("-")[0] == c.split("-")[0]] if t else []
        alt = [x for x in cands if x != c and fits(x, fam)]
        return (alt[0], t) if len(alt) == 1 else (None, t)

    node_links = []  # (nid, slot, how, conflict, map_code)
    for nid, n in nodes.items():
        # 同一節點同一年級只要有一個課號得靠校正清單才對得上，就整組都先試校正清單（同一批課號同一套編號）
        need = {c.split("-")[0] for c in n["courses"] if not (c in prof and fits(c, n["family"]))
                and translate(c, n["family"])[0]}
        for c in n["courses"]:
            alt, t = translate(c, n["family"])
            if c.split("-")[0] in need and alt:
                node_links.append((nid, alt, f"地圖課號 {c} 依 08-14 校正清單是「{t}」＝總表0816 {alt}，F 欄家族一致（同節點同年級整組換）", False, c))
            elif c in prof and fits(c, n["family"]):
                node_links.append((nid, c, "地圖課號＝總表0816課號，F 欄家族一致", False, c))
            else:
                node_links.append((nid, c, f"地圖課號 {c}＝總表0816「{prof[c]['title'] if c in prof else '？'}」，F 欄策略不屬本家族，校正清單也換不到", True, c))
    slot_links = collections.defaultdict(list)
    for nid, slot, how, conflict, mc in node_links:
        if (nid, slot) not in [(x[0], x[1]) for x in slot_links[slot]]:
            slot_links[slot].append((nid, slot, how, conflict, mc))

    def link_nodes(row):
        ls = [dict(id=nid, how=how, conflict=cf, map_code=mc, link_source="professor_ai_map")
              for nid, _, how, cf, mc in slot_links.get(row["slot"], [])]
        if any(not x["conflict"] for x in ls):
            return ls
        # 地圖沒列這課，或只列在家族不一致的節點 → 試規則：F 欄全文含某節點名，且該節點屬 F 欄家族（或 F 欄認不出家族）
        fs = families_of_strategy(row["strategy"])
        hits = [nid for nid, n in nodes.items() if row["strategy"] and norm(n["name"]) in norm(row["strategy"])
                and (not fs or n["family"] in fs)]
        if len(hits) == 1 and hits[0] not in [x["id"] for x in ls]:
            why = "教授地圖沒列這課" if not ls else "教授地圖只把這課放在家族不一致的節點"
            ls.append(dict(id=hits[0], how=f"{why}；F 欄「{row['strategy']}」含節點名「{nodes[hits[0]]['name']}」",
                           conflict=False, map_code=None, link_source="rule_from_professor_field"))
        return ls

    links = {sid: link_nodes(m["row"]) for sid, m in match.items()}
    stage_sets = collections.defaultdict(set)
    for nid, slot, how, cf, mc in node_links:
        if slot in prof:
            stage_sets[nid].add(STAGE_OF_GRADE[prof[slot]["grade"]])
    for sid, ls in links.items():
        for x in ls:
            stage_sets[x["id"]].add(STAGE_OF_GRADE[match[sid]["row"]["grade"]])
    node_stages = {nid: sorted(stage_sets[nid], key="ⅡⅢⅣ".index) for nid in nodes}
    resolved = {nid: resolve_node(sug, n, node_stages[nid]) for nid, n in nodes.items()}

    def node_layers(nid, stage):
        r = resolved[nid]
        perf = [code_ref(codes, c, source=r["source"], via=nid) for c in r["perf"][stage]]
        cont = [code_ref(codes, c, source=r["source"], via=nid) for c in r["content"][stage]]
        core = [code_ref(codes, f"國-{SCHOOL_OF_STAGE[stage]}-{r['core']}", source=r["source"], via=nid)]
        return perf, cont, core

    # ---- 3. 逐課 ----
    pending = []
    lessons_by_sid = {l["story_id"]: l for l in existing["lessons"]}
    node_online = collections.defaultdict(list)
    for s in stories:
        les = lessons_by_sid[s["id"]]
        m = match.get(s["id"])
        if not m:
            les["mapping_source"] = "ai_only"
            les["superseded_by_node"] = False
            continue
        row = m["row"]
        stage = STAGE_OF_GRADE[row["grade"]]
        les["professor_row"] = dict(
            source="professor_table", sheet=row["sheet"], excel_row=row["excel_row"], slot=row["slot"],
            title=row["title"], match=m["kind"], grade=row["grade"], genre=row["genre"], type=row["type"],
            issue=row["issue"], sixth_title=row["sixth_title"], text_type=row["text_type"], words=row["words"])
        les["professor_strategy"] = dict(text=row["strategy"], source="professor_table") if row["strategy"] else None
        if m["kind"] != "課號相同、課名相同":
            pending.append(dict(kind="課名差異", slot=row["slot"], story_id=s["id"],
                                detail=f"API「{s['title']}」／總表「{row['title']}」"))
        if les.get("stage") and les["stage"] != stage:
            pending.append(dict(kind="學習階段改由總表年級決定", slot=row["slot"], story_id=s["id"],
                                detail=f"舊草稿 {les['stage']} → 總表適用年級 {row['grade']} ＝第{stage}學習階段"))

        ls = links[s["id"]]
        for x in ls:
            if x["link_source"] == "rule_from_professor_field":
                pending.append(dict(kind="節點由規則補（教授地圖沒列這課）", slot=row["slot"], story_id=s["id"], detail=f"→ {x['id']}：{x['how']}"))
            elif x["map_code"] != row["slot"]:
                pending.append(dict(kind="教授地圖用舊課號，已依校正清單換到這課", slot=row["slot"], story_id=s["id"], detail=f"{x['id']} {nodes[x['id']]['name']}：{x['how']}"))
            elif x["conflict"]:
                pending.append(dict(kind="節點家族跟 F 欄策略不一致（不繼承代碼，請老師確認）", slot=row["slot"], story_id=s["id"], detail=f"{x['id']} {nodes[x['id']]['name']}（{fam_name[nodes[x['id']]['family']]}）：{x['how']}"))
        good = [x for x in ls if not x["conflict"]]
        inherit = good or ls
        if ls and not good:
            pending.append(dict(kind="這課只有家族不一致的節點（暫照地圖繼承）", slot=row["slot"], story_id=s["id"], detail="、".join(x["id"] for x in ls)))
        nids = [x["id"] for x in inherit]
        fam_hint = families_of_strategy(row["strategy"])
        primary = next((x for x in nids if nodes[x]["family"] in fam_hint), nids[0] if nids else None)
        les["professor_node"] = [dict(id=x["id"], name=nodes[x["id"]]["name"], family=nodes[x["id"]]["family"],
                                      family_name=fam_name[nodes[x["id"]]["family"]], primary=(x["id"] == primary),
                                      inherited=(x["id"] in nids), link_source=x["link_source"],
                                      map_code=x["map_code"], how=x["how"]) for x in ls]
        for x in nids:
            node_online[x].append(s["id"])

        perf, cont, core = [], [], []
        for x in nids:
            p, c, k = node_layers(x, stage)
            perf += p; cont += c; core += k
        gc, notes = genre_contents(codes, row["genre"], row["type"], stage)
        cont += gc
        if row["sheet"] == "3.文言文":
            cont += [code_ref(codes, c, source="rule_from_professor_field",
                              rule="R3 屬文言文分頁 → 常用文言文詞義、字詞虛字") for c in CLASSICAL_CODES]
        elif row["words"]:
            cont += [code_ref(codes, c, source="rule_from_professor_field",
                              rule=f"R2 語詞欄有 {len(row['words'])} 個 → 常用語詞認念／使用") for c in VOCAB_CODES[stage]]
        iss = issues_from(row["issue"], "議題") if row["sheet"] != "3.文言文" else issues_from(row["issue"], "單元議題")
        if row["sheet"] == "2.體育生的品格教材":
            iss += issues_from(row["strategy"], "F 欄策略")
        if row["issue"] and row["issue"] != "無" and not iss:
            notes.append(f"議題「{row['issue']}」沒有對應到 19 項議題")
        les["node_indicators"] = dict(
            stage=stage, stage_source=f"professor_table（{row['sheet']} 適用年級 {row['grade']}）",
            learning_performance=merge_items(perf), learning_content=merge_items(cont),
            core_competency=merge_items(core), issues=merge_items(iss), notes=notes)
        les["mapping_source"] = ("node_inherited" if good else "node_inherited_family_conflict") if nids else "ai_only_no_node"
        les["superseded_by_node"] = (dict(fields=["primary_indicator", "skills", "core_competencies", "learning_contents"],
                                          by=nids, note="AI 逐課草稿保留供對照；節點對照也還是建議，審核表回填前兩者都不是正式對照")
                                     if nids else False)

    # ---- 4. 頂層 ----
    existing["professor_source"] = dict(
        table=str(a.xlsx), table_note="陳淑麗教授 2026-08-18 傳送，說明「最後的總表，人工打的」；課綱代碼 0 筆",
        node_map=str(a.nodes),
        node_map_note="教授用 ChatGPT 做的課程地圖（chatgpt-log-1~4），我方萃取 90 節點；教授只覆核過推論家族，未封版 → 標 professor_ai_map，不標 professor_table",
        lesson_code_rule="總表（適用年級, 課次）＝二修課號＝API grade_code；體／文分頁用『新課次』",
        source_labels=dict(professor_table="教授總表儲存格直接有的值", rule_from_professor_field="從教授欄位用本腳本的固定規則推出",
                           professor_ai_map="教授以 ChatGPT 做的課程地圖（節點歸屬）", ai_inferred="Claude 推論（節點對課綱代碼）",
                           teacher_reviewed="老師在審核表確認後寫回"))
    existing["issues_19"] = dict(list=zg["issues_19"], source=zg["_meta"]["source_pdf"] + f"#page={zg['_meta']['issues_page'][0]}",
                                 note="十二年國教課程綱要總綱（111 學年度實施版）PDF 第 37 頁")
    used = sorted({it["code"] for l in existing["lessons"] if l.get("node_indicators")
                   for lay in ("learning_performance", "learning_content", "core_competency")
                   for it in l["node_indicators"][lay]})
    existing["node_codes"] = {c: dict(text=codes[c]["text"], pdf_page=codes[c]["k12ea_page"], url=codes[c]["k12ea_url"],
                                      naer_page=codes[c]["naer_page"], k12ea_copy=codes[c]["k12ea_copy"]) for c in used}
    existing["node_mapping"] = {
        nid: dict(name=n["name"], family=n["family"], family_name=fam_name[n["family"]], definition=n["definition"],
                  professor_courses=n["courses"],
                  resolved_courses=[dict(map_code=mc, slot=sl, how=how, conflict=cf) for x_nid, sl, how, cf, mc in node_links if x_nid == nid],
                  online_story_ids=sorted(set(node_online.get(nid, []))),
                  stages=node_stages[nid], suggestion=resolved[nid])
        for nid, n in nodes.items()}

    # ---- 5. 統計（程式算）----
    L = existing["lessons"]
    cnt = collections.Counter()
    for l in L:
        if l.get("professor_strategy"):
            cnt["professor_table"] += 1  # 策略
        if l.get("professor_row"):
            cnt["professor_table"] += 1  # 年級／學習階段
        for pn in l.get("professor_node") or []:
            if pn["inherited"]:
                cnt[pn["link_source"]] += 1
        ni = l.get("node_indicators") or {}
        for lay in ("learning_performance", "learning_content", "core_competency", "issues"):
            for it in ni.get(lay, []):
                cnt[it["source"]] += 1
    node_level = collections.Counter(r["source"] for r in resolved.values())
    cov = dict(
        computed_by="build_lesson_indicators.py", api_total=api["total"],
        lessons_matched_to_professor_table=f"{len(match)}/{len(L)}",
        lessons_with_professor_node=f"{sum(1 for l in L if l.get('professor_node'))}/{len(L)}",
        node_link_source=dict(collections.Counter(pn["link_source"] for l in L for pn in l.get("professor_node") or [] if pn["inherited"])),
        node_links_total=len(node_links),
        node_links_verified_0816=sum(1 for x in node_links if not x[3] and x[1] == x[4]),
        node_links_translated_from_old_code=sum(1 for x in node_links if not x[3] and x[1] != x[4]),
        node_links_family_conflict=sum(1 for x in node_links if x[3]),
        lessons_multi_node=sum(1 for l in L if len(l.get("professor_node") or []) > 1),
        three_layers_inherited_from_node=f"{sum(1 for l in L if str(l.get('mapping_source')).startswith('node_inherited'))}/{len(L)}",
        lessons_node_family_conflict_only=sum(1 for l in L if l.get("mapping_source") == "node_inherited_family_conflict"),
        lessons_three_layers_ai_only=sum(1 for l in L if l.get("mapping_source") in ("ai_only", "ai_only_no_node")),
        mapping_source_distribution=dict(collections.Counter(l.get("mapping_source") for l in L)),
        lesson_level_item_sources=dict(sorted(cnt.items(), key=lambda kv: SRC_RANK[kv[0]])),
        node_level_code_source=dict(node_level),
        nodes_total=len(nodes), nodes_with_online_lesson=sum(1 for n in nodes if node_online.get(n)),
        unmatched=unmatched, pending=pending,
        counting_rule="lesson_level_item_sources 一課一項目算一筆：策略 1＋年級 1（professor_table）、每個節點歸屬 1、三層每個代碼 1、每個議題 1；同課同層同代碼合併只算一次，來源取最強",
    )
    existing["node_mapping_coverage"] = cov
    rs = existing.get("review_status") or {}
    rs["node_mapping"] = ("2026-10-31 改為『教授節點對課綱』：課→節點來自教授課程地圖，節點→課綱代碼是 Claude 建議，"
                          "要老師審 90 節點（node-indicator-review.xlsx）回填後才是正式對照")
    existing["review_status"] = rs

    print(json.dumps({k: v for k, v in cov.items() if k not in ("pending", "unmatched")}, ensure_ascii=False, indent=1))
    print("unmatched:", json.dumps(unmatched, ensure_ascii=False))
    print("pending:", len(pending), dict(collections.Counter(p["kind"] for p in pending)))
    if a.check:
        return
    a.lesson_json.write_text(json.dumps(existing, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_review(a.review_xlsx, existing, nodes, fam_name, prof, sid_set, stories, cov)
    print("寫入", a.lesson_json, a.review_xlsx)


# ---------- 審核表 ----------
def write_review(path, data, nodes, fam_name, prof, sid_set, stories, cov):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "節點審核"
    online_slot = {api_slot(s["grade_code"]): s for s in stories}
    head = ["家族", "家族名", "節點", "節點名", "節點定義（教授地圖）", "教授的策略原文（總表 F 欄）",
            "屬於它的課（課號 課名）", "線上課數", "學習階段", "建議學習表現（代碼 原文）",
            "建議學習內容（代碼 原文）", "建議核心素養（代碼 原文）", "Claude 的理由", "逐課規則補強（依教授欄位）",
            "來源標記", "同意（打 V）", "修改為"]
    ws.append(head)
    lessons_by_sid = {l["story_id"]: l for l in data["lessons"]}

    def fmt(code):
        c = data["node_codes"].get(code)
        if c:
            return f"{code} {c['text']}"
        return code

    for nid, nm in data["node_mapping"].items():
        sugg = nm["suggestion"]
        courses = []
        strategies = []
        for rc in nm["resolved_courses"]:
            c = rc["slot"]
            row = prof.get(c)
            title = row["title"] if row else "（總表無）"
            mark = "" if c in online_slot else "（線上無）"
            if rc["conflict"]:
                mark += "【⚠ F 欄策略不屬本家族，未繼承】"
            elif rc["map_code"] != c:
                mark += f"【地圖原寫 {rc['map_code']}，依校正清單換】"
            courses.append(f"{c} {title}{mark}")
            if row and row["strategy"] and row["strategy"] not in strategies and not rc["conflict"]:
                strategies.append(row["strategy"])
        listed = {rc["slot"] for rc in nm["resolved_courses"]}
        for sid in nm["online_story_ids"]:  # 規則補進來的課
            pr = lessons_by_sid[sid].get("professor_row") or {}
            if pr.get("slot") not in listed:
                courses.append(f"{pr.get('slot')} {pr.get('title')}（規則補，教授地圖未列）")
                if pr and lessons_by_sid[sid]["professor_strategy"]:
                    t = lessons_by_sid[sid]["professor_strategy"]["text"]
                    if t not in strategies:
                        strategies.append(t)
        perf = "\n".join(f"[{st}] " + "；".join(fmt(c) for c in sugg["perf"][st]) for st in nm["stages"])
        cont = "\n".join(f"[{st}] " + "；".join(fmt(c) for c in sugg["content"][st]) for st in nm["stages"])
        core_codes = list(dict.fromkeys(f"國-{SCHOOL_OF_STAGE[st]}-{sugg['core']}" for st in nm["stages"]))
        core = "\n".join(fmt(c) for c in core_codes)
        rule_cnt = collections.Counter()
        for sid in nm["online_story_ids"]:
            ni = lessons_by_sid[sid].get("node_indicators") or {}
            for lay in ("learning_content", "issues"):
                for it in ni.get(lay, []):
                    if it["source"] == "rule_from_professor_field":
                        rule_cnt[it.get("code") or "議題:" + it["issue"]] += 1
        rules = "；".join(f"{k}×{v}課" for k, v in sorted(rule_cnt.items()))
        ws.append([nm["family"], nm["family_name"], nid, nm["name"], nm["definition"], "\n".join(strategies),
                   "\n".join(courses), len(nm["online_story_ids"]), "、".join(nm["stages"]), perf, cont, core,
                   sugg["why"], rules or "（無）",
                   f"課→節點：professor_ai_map；節點→代碼：{sugg['source']}；規則補強：rule_from_professor_field", "", ""])
    widths = [6, 12, 6, 18, 26, 26, 34, 7, 8, 46, 46, 40, 30, 30, 26, 10, 30]
    style(ws, widths, Alignment, Font, PatternFill, highlight_cols=(16, 17))

    ws2 = wb.create_sheet("逐課對照")
    ws2.append(["lesson_uid", "story_id", "API 課號", "API 課名", "總表列", "比對結果", "教授策略（F 欄）",
                "節點（主節點標★）", "節點來源", "學習階段", "學習表現", "學習內容", "核心素養", "議題（規則）", "備註"])
    for l in sorted(data["lessons"], key=lambda x: x["story_id"]):
        pr = l.get("professor_row") or {}
        ni = l.get("node_indicators") or {}
        def j(lay):
            return "\n".join(f"{it['code']}（{it['source']}）" for it in ni.get(lay, []))
        ws2.append([l["lesson_uid"], l["story_id"], l["grade_code"], l["title"],
                    f"{pr.get('sheet')} 第{pr.get('excel_row')}列" if pr else "（對不上）", pr.get("match", ""),
                    (l.get("professor_strategy") or {}).get("text", ""),
                    "\n".join(("★" if n["primary"] else "") + f"{n['id']} {n['name']}" for n in l.get("professor_node") or []),
                    "、".join(sorted({n["link_source"] for n in l.get("professor_node") or []})),
                    ni.get("stage", ""), j("learning_performance"), j("learning_content"), j("core_competency"),
                    "、".join(it["issue"] for it in ni.get("issues", [])), "\n".join(ni.get("notes", []))])
    style(ws2, [8, 8, 9, 26, 16, 22, 28, 28, 18, 6, 30, 34, 24, 14, 30], Alignment, Font, PatternFill)

    ws3 = wb.create_sheet("對不上與待確認")
    ws3.append(["類別", "課號", "story_id", "說明"])
    for u in cov["unmatched"]:
        ws3.append(["對不上總表", u["grade_code"], u["story_id"], f"{u['title']}：{u['reason']}"])
    for p in cov["pending"]:
        ws3.append([p["kind"], p["slot"], p["story_id"], p["detail"]])
    style(ws3, [26, 10, 10, 90], Alignment, Font, PatternFill)

    ws4 = wb.create_sheet("說明與統計")
    for line in [
        ["這份表在做什麼", "把 108 課綱三層對到教授的 90 個節點；每課繼承它所屬節點的代碼，所以老師只審 90 列不審 179 課"],
        ["審核狀態", "全部是建議，老師在「節點審核」最後兩欄回填前，不得當正式對照"],
        ["課綱條文", "只抄教育部官方 PDF 原文；頁碼是 k12ea PDF 檔第幾頁（國教院版 +2，程式驗過 221 條一致）"],
        ["來源標記", "professor_table＝教授總表直接有；rule_from_professor_field＝從教授欄位用固定規則推；professor_ai_map＝教授用 ChatGPT 做的課程地圖（未封版）；ai_inferred＝Claude 推論"],
        ["179 課對到總表", cov["lessons_matched_to_professor_table"]],
        ["對到教授節點", cov["lessons_with_professor_node"] + "　" + json.dumps(cov["node_link_source"], ensure_ascii=False)],
        ["逐課項目來源統計", json.dumps(cov["lesson_level_item_sources"], ensure_ascii=False)],
        ["節點層代碼來源", json.dumps(cov["node_level_code_source"], ensure_ascii=False)],
        ["議題清單來源", data["issues_19"]["source"]],
        ["重跑", "python3 .claude/skills/lesson-curriculum-mapping/scripts/build_lesson_indicators.py"],
    ]:
        ws4.append(line)
    style(ws4, [20, 110], Alignment, Font, PatternFill)
    wb.save(path)


def style(ws, widths, Alignment, Font, PatternFill, highlight_cols=()):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in highlight_cols:
        for r in range(1, ws.max_row + 1):
            ws.cell(r, col).fill = PatternFill("solid", fgColor="FFF2CC")


if __name__ == "__main__":
    sys.exit(main())
